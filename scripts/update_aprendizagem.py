from __future__ import annotations
import io, json, re, unicodedata, urllib.request
from pathlib import Path
from openpyxl import load_workbook

SOURCE = "https://www.gov.br/trabalho-e-emprego/pt-br/assuntos/inspecao-do-trabalho/areas-de-atuacao/insercao-de-aprendiz-1/potencial_ead_julho-1.xlsx"
IBGE = "https://servicodados.ibge.gov.br/api/v1/localidades/municipios?orderBy=nome"
OUT = Path("data/snapshots/aprendizagem.json")

def norm(v):
    s = "" if v is None else str(v)
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()

def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent":"RedeCidada-Snapshot/1.0"})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()

cities = json.loads(fetch(IBGE).decode("utf-8"))
lookup = {}
for x in cities:
    uf = (((x.get("regiao-imediata") or {}).get("regiao-intermediaria") or {}).get("UF") or {}).get("sigla") or (((x.get("microrregiao") or {}).get("mesorregiao") or {}).get("UF") or {}).get("sigla") or ""
    lookup[(norm(uf), norm(x.get("nome")))] = str(x["id"])

wb = load_workbook(io.BytesIO(fetch(SOURCE)), read_only=True, data_only=True)
out = {}
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        vals = [norm(v) for v in row]
        if any(v in {"uf","estado"} for v in vals) and any("municip" in v for v in vals):
            header = [str(v or "").strip() for v in row]
            break
    if not header:
        continue
    nh = [norm(v) for v in header]
    def idx(*tests):
        for i,h in enumerate(nh):
            if any(t in h for t in tests):
                return i
        return -1
    iuf = idx("uf","estado")
    im = idx("municip")
    itotal = idx("potencial","total")
    for row in rows:
        if iuf < 0 or im < 0 or len(row) <= max(iuf,im):
            continue
        uf, city = norm(row[iuf]), norm(row[im])
        code = lookup.get((uf,city))
        if not code:
            continue
        nums = []
        for j,v in enumerate(row):
            if j in (iuf,im):
                continue
            try:
                if v is None or str(v).strip()=="":
                    continue
                n = float(str(v).replace(".","").replace(",","."))
                nums.append((j,n))
            except Exception:
                pass
        if not nums:
            continue
        total = None
        if itotal >= 0:
            try:
                total = float(str(row[itotal]).replace(".","").replace(",","."))
            except Exception:
                pass
        if total is None:
            total = nums[0][1]
        segments = {}
        for j,n in nums:
            if j == itotal or j >= len(header):
                continue
            label = header[j].strip()
            if label and norm(label) not in {"total","potencial","potencial total"}:
                segments[label] = int(n) if n.is_integer() else n
        out[code] = {
            "total": int(total) if float(total).is_integer() else total,
            "potential": int(total) if float(total).is_integer() else total,
            "segments": segments,
            "reference": "jul/2026",
            "source": "MTE/SIT · eSocial"
        }

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",",":")), encoding="utf-8")
print(f"{len(out)} municípios gravados em {OUT}")
